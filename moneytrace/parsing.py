"""In-memory parsing for CSV, XLSX and text-based PDF bank statements.

Bank formats vary. Never silently invent a debit/credit direction for ambiguous rows.
"""
from __future__ import annotations

from dataclasses import dataclass, field, asdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
import csv
import hashlib
import re

from dateutil import parser as date_parser
from openpyxl import load_workbook
from pypdf import PdfReader

MAX_FILE_BYTES = 6 * 1024 * 1024
MAX_RECORDS = 20000
MAX_PDF_PAGES = 60


class StatementError(ValueError):
    pass


@dataclass
class Transaction:
    id: str
    date: str
    description: str
    debit: Decimal
    credit: Decimal
    balance: Decimal | None
    reference: str
    account: str
    source: str
    line: int
    confidence: str = "high"

    def serialize(self):
        d = asdict(self)
        for key in ("debit", "credit", "balance"):
            if d[key] is not None:
                d[key] = f"{d[key]:.2f}"
        return d


@dataclass
class Statement:
    filename: str
    account: str
    sha256: str
    received_at: str
    format: str
    transactions: list[Transaction] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    skipped_rows: int = 0

    def serialize(self):
        return {
            "filename": self.filename,
            "account": self.account,
            "sha256": self.sha256,
            "received_at": self.received_at,
            "format": self.format,
            "transaction_count": len(self.transactions),
            "skipped_rows": self.skipped_rows,
            "warnings": self.warnings,
        }


ALIASES = {
    "date": ["date", "transactiondate", "value date", "valuedate", "postingdate", "transdate", "trans date"],
    "description": ["description", "narration", "details", "transactiondetails", "transaction description", "remarks", "particulars", "memo"],
    "debit": ["debit", "dr", "withdrawal", "withdrawals", "moneyout", "paidout", "outflow"],
    "credit": ["credit", "cr", "deposit", "deposits", "moneyin", "paidin", "inflow"],
    "balance": ["balance", "runningbalance", "closingbalance", "availablebalance", "ledgerbalance"],
    "amount": ["amount", "transactionamount", "value"],
    "type": ["type", "transactiontype", "direction", "drcr", "entrytype"],
    "reference": ["reference", "ref", "transactionreference", "sessionid", "transactionid", "rrn", "paymentreference"],
    "status": ["status", "transactionstatus"],
}


def normalize_heading(raw) -> str:
    return re.sub(r"[^a-z0-9]", "", str(raw or "").lower())


def identify_columns(headers):
    found = {}
    for idx, value in enumerate(headers):
        h = normalize_heading(value)
        # Some bank exports include units and column adornments, e.g. Debit (NGN).
        h = re.sub(r"(ngn|naira|wat)$", "", h)
        if h.endswith("amount") and h not in ("amount", "transactionamount"):
            h = h[:-6]
        if not h:
            continue
        for key, synonyms in ALIASES.items():
            if key not in found and h in {normalize_heading(s) for s in synonyms}:
                found[key] = idx
    if "date" not in found or not ("debit" in found or "credit" in found or "amount" in found):
        return None
    return found


def number(value) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        n = Decimal(str(value))
        return n.quantize(Decimal("0.01")) if n.is_finite() else None
    text = str(value).strip()
    if not text or text.lower() in ("-", "n/a", "na", "null", "none"):
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"(?i)(NGN|₦|\s|,)", "", text).strip("()")
    text = re.sub(r"(?i)(DR|CR)$", "", text)
    if not re.fullmatch(r"[-+]?\d+(?:\.\d{1,2})?", text):
        return None
    try:
        result = Decimal(text).quantize(Decimal("0.01"))
        return -abs(result) if negative else result
    except InvalidOperation:
        return None


def parse_date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    if not text or len(text) < 6 or len(text) > 45:
        return None
    try:
        d = date_parser.parse(text, dayfirst=not bool(re.match(r"^\d{4}[-/]", text)), fuzzy=False)
        if d.year < 2000 or d.year > 2100:
            return None
        return d.date().isoformat()
    except (ValueError, OverflowError, TypeError):
        return None


def account_from_filename(filename):
    # A user-provided label, not proof of bank identity.
    name = re.sub(r"[_-]+", " ", Path(filename).stem).strip()
    return (name[:65] or "Uploaded statement")


def parse_rows(rows, filename, fmt, digest, timestamp) -> Statement:
    stmt = Statement(filename, account_from_filename(filename), digest, timestamp, fmt)
    header_idx = -1
    cols = None
    for idx, row in enumerate(rows[:35]):
        cols = identify_columns(row)
        if cols:
            header_idx = idx
            break
    if not cols:
        raise StatementError("Could not identify date and debit/credit columns. Export a CSV statement with Date, Description, Debit, Credit and Balance headers.")

    def cell(row, key):
        col = cols.get(key)
        return row[col] if col is not None and col < len(row) else None

    for rownum, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if len(stmt.transactions) >= MAX_RECORDS:
            stmt.warnings.append(f"Stopped at the {MAX_RECORDS:,}-transaction safety limit.")
            break
        if not any(str(c or "").strip() for c in row):
            continue
        raw_date = cell(row, "date")
        d = parse_date(raw_date)
        if not d:
            stmt.skipped_rows += 1
            continue
        debit = number(cell(row, "debit"))
        credit = number(cell(row, "credit"))
        amount = number(cell(row, "amount"))
        tx_type = str(cell(row, "type") or "").lower().strip()
        if debit is None and credit is None and amount is not None:
            if any(t in tx_type for t in ("debit", "withdraw", "dr", "out")):
                debit, credit = abs(amount), Decimal("0")
            elif any(t in tx_type for t in ("credit", "deposit", "cr", "in")):
                debit, credit = Decimal("0"), abs(amount)
            elif amount < 0:
                debit, credit = abs(amount), Decimal("0")
            else:
                # A positive amount without a direction is ambiguous.
                stmt.skipped_rows += 1
                continue
        if debit is None and credit is None:
            stmt.skipped_rows += 1
            continue
        if debit is not None and debit < 0 and credit is None:
            debit = abs(debit)
        debit = debit or Decimal("0")
        credit = credit or Decimal("0")
        if debit < 0 or credit < 0 or (debit and credit):
            stmt.skipped_rows += 1
            continue
        desc = str(cell(row, "description") or "Transaction")[:320].strip()
        status = str(cell(row, "status") or "").strip()
        if status and status.lower() not in desc.lower():
            desc += f" [{status[:25]}]"
        ref = str(cell(row, "reference") or "").strip()[:75]
        bal = number(cell(row, "balance"))
        stmt.transactions.append(Transaction(
            id=f"{digest[:9]}-{rownum}", date=d, description=desc,
            debit=debit, credit=credit, balance=bal, reference=ref,
            account=stmt.account, source=filename, line=rownum,
        ))
    if stmt.skipped_rows:
        stmt.warnings.append(f"{stmt.skipped_rows} rows could not be interpreted as unambiguous transactions; review the original statement.")
    if not stmt.transactions:
        raise StatementError("No unambiguous transaction rows were found. Check the date and debit/credit column format.")
    if not any(t.balance is not None for t in stmt.transactions):
        stmt.warnings.append("No running-balance column found; balance-gap detection is unavailable for this statement.")
    return stmt


def parse_pdf(raw, filename, digest, timestamp):
    try:
        pdf = PdfReader(BytesIO(raw), strict=False)
        if pdf.is_encrypted:
            raise StatementError("Encrypted PDFs are not supported. Export an unlocked statement from your bank.")
        if len(pdf.pages) > MAX_PDF_PAGES:
            raise StatementError(f"PDFs are limited to {MAX_PDF_PAGES} pages.")
        blocks = []
        for page in pdf.pages:
            blocks.append((page.extract_text(extraction_mode="layout") or "")[:300000])
        content = "\n".join(blocks)
    except StatementError:
        raise
    except Exception:
        raise StatementError("The PDF could not be read. Please export a CSV or XLSX statement.")
    if len(content.strip()) < 60:
        raise StatementError("This PDF appears scanned or image-only. Text extraction is not available; export CSV/XLSX instead. No OCR was performed.")
    # Conservative parser for text PDFs with a single row per transaction and trailing
    # debit, credit, balance fields. Ambiguous two-amount rows are intentionally skipped.
    date_pattern = re.compile(r"^\s*(\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{2,4})\s+(.+?)\s*$")
    amount_pattern = re.compile(r"[-+]?(?:₦\s*)?\d{1,3}(?:,\d{3})+(?:\.\d{2})?|[-+]?(?:₦\s*)?\d+\.\d{2}")
    stmt = Statement(filename, account_from_filename(filename), digest, timestamp, "pdf")
    candidate = 0
    for lineno, line in enumerate(content.splitlines(), start=1):
        match = date_pattern.match(line)
        if not match:
            continue
        d = parse_date(match.group(1))
        if not d:
            continue
        candidate += 1
        tail = match.group(2)
        entries = list(amount_pattern.finditer(tail))
        if len(entries) < 3:
            stmt.skipped_rows += 1
            continue
        last3 = entries[-3:]
        suffix = tail[last3[-1].end():].strip()
        # Reject table rows where trailing fields aren't clearly amounts.
        if suffix and not re.fullmatch(r"(?:CR|DR|NGN|₦|\s)*", suffix, re.IGNORECASE):
            stmt.skipped_rows += 1
            continue
        debit, credit, balance = [number(a.group(0)) for a in last3]
        if None in (debit, credit, balance) or debit < 0 or credit < 0 or (debit and credit):
            stmt.skipped_rows += 1
            continue
        desc = tail[:last3[0].start()].strip()
        if len(desc) < 2:
            stmt.skipped_rows += 1
            continue
        ref_match = re.search(r"\b(?:REF|RRN|SESSION)[:# -]*([A-Za-z0-9-]{6,})", desc, re.I)
        stmt.transactions.append(Transaction(
            id=f"{digest[:9]}-{lineno}", date=d, description=desc[:320],
            debit=debit, credit=credit, balance=balance,
            reference=ref_match.group(1) if ref_match else "", account=stmt.account,
            source=filename, line=lineno, confidence="review_needed",
        ))
    if not stmt.transactions:
        raise StatementError("The PDF text layout is not safely parseable. Export CSV or XLSX for a reliable audit; no transactions were guessed.")
    stmt.warnings.append("PDF rows were interpreted from extracted text (assumed columns: debit, credit, balance). Confirm the column order and amounts against your original statement before relying on findings.")
    if stmt.skipped_rows:
        stmt.warnings.append(f"{stmt.skipped_rows} PDF rows looked like dated entries but could not be parsed safely.")
    return stmt


def parse_file(raw: bytes, filename: str, timestamp: str) -> Statement:
    if len(raw) > MAX_FILE_BYTES:
        raise StatementError("Each file must be 6 MB or smaller.")
    if not raw:
        raise StatementError("File is empty.")
    filename = Path(filename or "statement").name[:110]
    ext = Path(filename).suffix.lower()
    digest = hashlib.sha256(raw).hexdigest()
    if ext == ".csv":
        try:
            txt = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            txt = raw.decode("cp1252", errors="replace")
        if "\x00" in txt:
            raise StatementError("Invalid CSV content.")
        try:
            dialect = csv.Sniffer().sniff(txt[:4000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(StringIO(txt), dialect=dialect))[:MAX_RECORDS + 40]
        return parse_rows(rows, filename, "csv", digest, timestamp)
    if ext == ".xlsx":
        try:
            wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
            errors = []
            for sheet in wb.worksheets[:4]:
                rows = list(sheet.iter_rows(max_row=MAX_RECORDS + 40, values_only=True))
                try:
                    return parse_rows(rows, filename, "xlsx", digest, timestamp)
                except StatementError as exc:
                    errors.append(str(exc))
        except Exception:
            raise StatementError("XLSX workbook could not be read.")
        raise StatementError("No worksheet had usable transaction columns.")
    if ext == ".pdf":
        return parse_pdf(raw, filename, digest, timestamp)
    raise StatementError("Unsupported format. Use CSV, XLSX, or a text-based PDF.")
